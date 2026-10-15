import os
from datetime import datetime
from typing import List, Dict, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.config import EXPORTS_DIR, CURRENCY, SELLER_DEFAULT_NAME

# Palette styling
NAVY_HEADER = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
GOLD_ACCENT = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
GREEN_STATUS = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
YELLOW_STATUS = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
RED_STATUS = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
ZEBRA_ROW = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="1E293B")
FONT_SUBTITLE = Font(name="Calibri", size=11, italic=True, color="64748B")
FONT_REGULAR = Font(name="Calibri", size=10, color="0F172A")
FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="0F172A")
FONT_STRIKE = Font(name="Calibri", size=10, strike=True, color="94A3B8")

THIN_BORDER_SIDE = Side(border_style="thin", color="E2E8F0")
BORDER_CELL = Border(top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE, left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE)

def auto_fit_columns(ws, min_width=12):
    """Automatically adjusts column widths for clean readability."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip title merged banner rows in width calculation
            if cell.row in [1, 2, 3]:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, min_width)

def export_production_sheet(orders: List[Dict[str, Any]], target_date: Optional[str] = None, seller_name: str = SELLER_DEFAULT_NAME) -> str:
    """
    Generates the Baker's Morning Kitchen Production Sheet in Excel.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Kitchen Production"
    ws.views.sheetView[0].showGridLines = True

    if not target_date:
        target_date = datetime.now().strftime("%Y-%m-%d")

    # Title & Metadata
    ws.merge_cells("A1:I1")
    ws["A1"] = f"{seller_name.upper()} — DAILY KITCHEN PRODUCTION SHEET"
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells("A2:I2")
    ws["A2"] = f"Production Schedule for: {target_date}  |  Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = FONT_SUBTITLE

    # Headers
    headers = [
        "Order #", "Delivery Time", "Customer", "Product", "Size / Variant",
        "Qty", "Custom Inscription / Kitchen Notes", "Status", "Kitchen Done"
    ]
    row_num = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=header)
        cell.font = FONT_HEADER
        cell.fill = NAVY_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_CELL
    ws.row_dimensions[row_num].height = 26

    # Populate order items
    total_items_count = 0
    row_num = 5

    for order in orders:
        items = order.get("items", [])
        status = order.get("status", "new").lower()

        for it in items:
            total_items_count += it.get("quantity", 1)

            ws.cell(row=row_num, column=1, value=f"#{order.get('id', '')}").alignment = Alignment(horizontal="center")
            ws.cell(row=row_num, column=2, value=order.get("delivery_time", "Anytime"))
            ws.cell(row=row_num, column=3, value=order.get("customer_name", "Customer"))
            ws.cell(row=row_num, column=4, value=it.get("product_name", ""))
            ws.cell(row=row_num, column=5, value=it.get("variant", ""))
            ws.cell(row=row_num, column=6, value=it.get("quantity", 1)).alignment = Alignment(horizontal="center")

            notes_cell = ws.cell(row=row_num, column=7, value=it.get("notes") or order.get("user_notes") or "-")
            if it.get("notes"):
                notes_cell.fill = GOLD_ACCENT
                notes_cell.font = FONT_BOLD

            status_cell = ws.cell(row=row_num, column=8, value=status.upper())
            status_cell.alignment = Alignment(horizontal="center")

            if status == "cancelled":
                status_cell.fill = RED_STATUS
                for col in range(1, 10):
                    ws.cell(row=row_num, column=col).font = FONT_STRIKE
            elif status == "changed":
                status_cell.fill = YELLOW_STATUS
            else:
                status_cell.fill = GREEN_STATUS

            done_cell = ws.cell(row=row_num, column=9, value="[   ]")
            done_cell.alignment = Alignment(horizontal="center")

            # Apply borders and styling
            for col in range(1, 10):
                c = ws.cell(row=row_num, column=col)
                c.border = BORDER_CELL
                if status != "cancelled" and not (col == 7 and it.get("notes")):
                    if row_num % 2 == 0:
                        c.fill = ZEBRA_ROW
                    c.font = FONT_REGULAR

            ws.row_dimensions[row_num].height = 22
            row_num += 1

    # Summary section
    row_num += 1
    ws.merge_cells(f"A{row_num}:E{row_num}")
    summary_label = ws.cell(row=row_num, column=1, value="TOTAL PRODUCTION ITEMS:")
    summary_label.font = FONT_BOLD
    summary_label.alignment = Alignment(horizontal="right")

    sum_cell = ws.cell(row=row_num, column=6, value=total_items_count)
    sum_cell.font = FONT_BOLD
    sum_cell.alignment = Alignment(horizontal="center")
    sum_cell.border = BORDER_CELL

    auto_fit_columns(ws)

    filename = f"Production_Sheet_{target_date}.xlsx"
    filepath = EXPORTS_DIR / filename
    wb.save(filepath)
    return str(filepath)

def export_delivery_sheet(orders: List[Dict[str, Any]], target_date: Optional[str] = None, seller_name: str = SELLER_DEFAULT_NAME) -> str:
    """
    Generates Dispatch & Delivery Sheet grouped by Area.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Deliveries by Area"
    ws.views.sheetView[0].showGridLines = True

    if not target_date:
        target_date = datetime.now().strftime("%Y-%m-%d")

    ws.merge_cells("A1:I1")
    ws["A1"] = f"{seller_name.upper()} — DISPATCH & DELIVERY SHEET BY AREA"
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells("A2:I2")
    ws["A2"] = f"Delivery Route Date: {target_date}  |  Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = FONT_SUBTITLE

    headers = [
        "Area / Route", "Time Slot", "Customer Name", "Phone Number",
        "Delivery Address / Notes", "Items Summary", "Total (LKR)", "Payment Status", "Recipient Signature"
    ]
    row_num = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=header)
        cell.font = FONT_HEADER
        cell.fill = NAVY_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_CELL
    ws.row_dimensions[row_num].height = 26

    # Sort orders by Area
    sorted_orders = sorted(orders, key=lambda x: (x.get("area") or "Unknown Area", x.get("delivery_time") or ""))

    row_num = 5
    total_cod_amount = 0.0

    for order in sorted_orders:
        if order.get("status") == "cancelled":
            continue

        items_summary = ", ".join(f"{int(it.get('quantity', 1))}x {it.get('product_name')} ({it.get('variant')})" for it in order.get("items", []))
        total_amt = float(order.get("total_amount", 0.0))
        pay_status = order.get("payment_status", "pending").lower()

        if pay_status in ["cod", "pending"]:
            total_cod_amount += total_amt
            pay_display = f"COLLECT {CURRENCY} {total_amt:,.2f}"
            pay_fill = YELLOW_STATUS
        else:
            pay_display = "PAID (Bank Transfer)"
            pay_fill = GREEN_STATUS

        ws.cell(row=row_num, column=1, value=order.get("area") or "Local Pickup")
        ws.cell(row=row_num, column=2, value=order.get("delivery_time", "Anytime"))
        ws.cell(row=row_num, column=3, value=order.get("customer_name", "Customer"))
        ws.cell(row=row_num, column=4, value=order.get("customer_phone", "-"))
        ws.cell(row=row_num, column=5, value=order.get("delivery_address") or order.get("area") or "-")
        ws.cell(row=row_num, column=6, value=items_summary)

        amt_cell = ws.cell(row=row_num, column=7, value=total_amt)
        amt_cell.number_format = "#,##0.00"
        amt_cell.alignment = Alignment(horizontal="right")

        p_cell = ws.cell(row=row_num, column=8, value=pay_display)
        p_cell.fill = pay_fill
        p_cell.alignment = Alignment(horizontal="center")
        p_cell.font = FONT_BOLD

        ws.cell(row=row_num, column=9, value="")

        for col in range(1, 10):
            c = ws.cell(row=row_num, column=col)
            c.border = BORDER_CELL
            if col != 8:
                c.font = FONT_REGULAR

        ws.row_dimensions[row_num].height = 24
        row_num += 1

    # Total Cash to Collect
    row_num += 1
    ws.merge_cells(f"A{row_num}:F{row_num}")
    lbl = ws.cell(row=row_num, column=1, value="TOTAL CASH TO COLLECT (COD / PENDING):")
    lbl.font = FONT_BOLD
    lbl.alignment = Alignment(horizontal="right")

    tot_cell = ws.cell(row=row_num, column=7, value=total_cod_amount)
    tot_cell.number_format = f'"{CURRENCY} " #,##0.00'
    tot_cell.font = FONT_BOLD
    tot_cell.fill = GOLD_ACCENT
    tot_cell.border = BORDER_CELL
    tot_cell.alignment = Alignment(horizontal="right")

    auto_fit_columns(ws)

    filename = f"Delivery_Sheet_{target_date}.xlsx"
    filepath = EXPORTS_DIR / filename
    wb.save(filepath)
    return str(filepath)

def export_payment_pending_sheet(orders: List[Dict[str, Any]], seller_name: str = SELLER_DEFAULT_NAME) -> str:
    """
    Generates Payment-Pending and Cash-on-Delivery Ledger.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Pending Payments"
    ws.views.sheetView[0].showGridLines = True

    ws.merge_cells("A1:H1")
    ws["A1"] = f"{seller_name.upper()} — PENDING PAYMENTS & ADVANCES TRACKER"
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Report Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = FONT_SUBTITLE

    headers = [
        "Order #", "Delivery Date", "Customer Name", "Phone Number",
        "Items Summary", "Total Amount (LKR)", "Payment Type", "Follow-up Status"
    ]
    row_num = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=header)
        cell.font = FONT_HEADER
        cell.fill = NAVY_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_CELL
    ws.row_dimensions[row_num].height = 26

    row_num = 5
    total_pending = 0.0

    for order in orders:
        if order.get("status") == "cancelled":
            continue
        pay_status = order.get("payment_status", "pending").lower()
        if pay_status == "paid":
            continue

        total_amt = float(order.get("total_amount", 0.0))
        total_pending += total_amt
        items_summary = ", ".join(f"{int(it.get('quantity', 1))}x {it.get('product_name')}" for it in order.get("items", []))

        ws.cell(row=row_num, column=1, value=f"#{order.get('id', '')}").alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=2, value=order.get("delivery_date", "-")).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=3, value=order.get("customer_name", "Customer"))
        ws.cell(row=row_num, column=4, value=order.get("customer_phone", "-"))
        ws.cell(row=row_num, column=5, value=items_summary)

        amt_cell = ws.cell(row=row_num, column=6, value=total_amt)
        amt_cell.number_format = "#,##0.00"
        amt_cell.alignment = Alignment(horizontal="right")

        ws.cell(row=row_num, column=7, value="Cash on Delivery" if pay_status == "cod" else "Pending Slip").alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=8, value="Send WhatsApp reminder" if pay_status == "pending" else "Collect at door")

        for col in range(1, 9):
            c = ws.cell(row=row_num, column=col)
            c.border = BORDER_CELL
            c.font = FONT_REGULAR

        ws.row_dimensions[row_num].height = 22
        row_num += 1

    # Total row
    row_num += 1
    ws.merge_cells(f"A{row_num}:E{row_num}")
    lbl = ws.cell(row=row_num, column=1, value="TOTAL PENDING REVENUE:")
    lbl.font = FONT_BOLD
    lbl.alignment = Alignment(horizontal="right")

    tot_cell = ws.cell(row=row_num, column=6, value=total_pending)
    tot_cell.number_format = f'"{CURRENCY} " #,##0.00'
    tot_cell.font = FONT_BOLD
    tot_cell.fill = GOLD_ACCENT
    tot_cell.border = BORDER_CELL
    tot_cell.alignment = Alignment(horizontal="right")

    auto_fit_columns(ws)

    filename = f"Payment_Pending_Sheet_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    filepath = EXPORTS_DIR / filename
    wb.save(filepath)
    return str(filepath)
